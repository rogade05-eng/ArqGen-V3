"""Catálogo PRECONS III real (v1.9.0): ruleset versionado RoPres 3.30.

Contratos: 15981 renglones y 4383 recursos oficiales como datos, 38
secciones, desglose oficial (materiales+mano de obra+equipos = total),
parámetros y límites normativos, índices O(1), búsqueda sin acentos,
servicio de consulta, CLI (list filtrado + precio) y navegador GUI.
"""

from __future__ import annotations

import os
import unittest

from core.errors import DomainError
from tests.base import ARQGenTestCase

RULESET = "precons_iii_rospres_330"
EXPECTED_ITEMS = 15981
EXPECTED_RESOURCES = 4383


def _ruleset():
    from app.paths import resource_path
    from engines.precons_engine import PreconsRuleset
    return PreconsRuleset.from_file(
        resource_path("rulesets", f"{RULESET}.json"))


class TestCatalogoData(ARQGenTestCase):
    """Integridad del catálogo completo como datos (spec 59)."""

    def test_load_full_catalog(self):
        ruleset = _ruleset()
        self.assertEqual(ruleset.code, RULESET)
        self.assertEqual(ruleset.version, "3.30")
        self.assertEqual(ruleset.currency, "CUP")
        self.assertEqual(len(ruleset.work_items), EXPECTED_ITEMS)
        self.assertEqual(len(ruleset.resources), EXPECTED_RESOURCES)

    def test_unique_codes(self):
        ruleset = _ruleset()
        item_codes = [i["code"] for i in ruleset.work_items]
        resource_codes = [r["code"] for r in ruleset.resources]
        self.assertEqual(len(set(item_codes)), EXPECTED_ITEMS)
        self.assertEqual(len(set(resource_codes)), EXPECTED_RESOURCES)

    def test_resource_kinds_official(self):
        ruleset = _ruleset()
        kinds = {r["kind"] for r in ruleset.resources}
        self.assertEqual(kinds, {"MATERIAL", "EQUIPMENT", "LABOR"})
        for resource in ruleset.resources:
            self.assertTrue(resource.get("unit"), resource["code"])
            self.assertGreaterEqual(float(resource.get("price_ref", 0)), 0)

    def test_items_have_unit_and_positive_official_price(self):
        ruleset = _ruleset()
        for item in ruleset.work_items:
            self.assertTrue(item.get("unit"), item["code"])
            self.assertGreater(float(item.get("price_ref", 0)), 0,
                               item["code"])
            self.assertTrue(str(item.get("sec", "")), item["code"])

    def test_official_breakdown_consistent(self):
        ruleset = _ruleset()
        for item in ruleset.work_items:
            total = (float(item["price_material"])
                     + float(item["price_labor"])
                     + float(item["price_equipment"]))
            self.assertAlmostEqual(total, float(item["price_ref"]),
                                   places=2, msg=item["code"])

    def test_parameters_limits_and_source(self):
        ruleset = _ruleset()
        self.assertEqual(len(ruleset.parameters), 8)
        self.assertEqual(ruleset.parameters["acarreo_horizontal_m"]["value"],
                         40)
        self.assertEqual(ruleset.parameters["horas_jornada"]["value"], 8)
        self.assertEqual(len(ruleset.limits), 4)
        self.assertEqual(ruleset.limits["FACTOR_TARIFARIO_MAX"]["value"], 5)
        self.assertEqual(ruleset.limits["UTILIDAD_MAX"]["value"], 0.15)
        self.assertEqual(ruleset.source["rospres"], "3.30")
        self.assertEqual(ruleset.generated, "2026-09-24")
        self.assertEqual(len(ruleset.source["instrucciones"]), 7)
        self.assertEqual(ruleset.variants[0]["code"], "BASE")


class TestCatalogoEngine(ARQGenTestCase):
    """Acceso O(1), secciones, búsqueda y precio oficial."""

    def test_known_item_lookup(self):
        from engines.precons_engine import official_price
        ruleset = _ruleset()
        item = ruleset.item("011111")
        self.assertEqual(item["name"], "MANIGUA MANUAL")
        self.assertEqual(item["unit"], "m2")
        price = official_price(ruleset, "011111")
        self.assertAlmostEqual(price["total"], 26.81, places=2)
        self.assertAlmostEqual(price["material"], 26.81, places=2)
        self.assertAlmostEqual(price["labor"], 0.0, places=2)
        self.assertEqual(price["currency"], "CUP")

    def test_unknown_item_error_is_capped(self):
        ruleset = _ruleset()
        with self.assertRaises(DomainError) as ctx_manager:
            ruleset.item("XX-NO-EXISTE")
        context = ctx_manager.exception.context
        self.assertEqual(context["total"], EXPECTED_ITEMS)
        self.assertLessEqual(len(context["available"]), 15)

    def test_sections_index(self):
        ruleset = _ruleset()
        sections = ruleset.sections()
        self.assertEqual(len(sections), 38)
        first = sections[0]
        self.assertEqual(first["code"], "01")
        self.assertEqual(first["items"], 653)
        self.assertIn("Movimiento", first["name"])

    def test_search_case_and_accent_insensitive(self):
        ruleset = _ruleset()
        hits = ruleset.search("HORMIGÓN", section="04", limit=200)
        self.assertTrue(hits)
        for item in hits:
            self.assertEqual(item["sec"], "04")
        plain = ruleset.search("hormigon", section="04", limit=200)
        self.assertEqual([h["code"] for h in hits],
                         [p["code"] for p in plain])

    def test_search_by_section_only(self):
        ruleset = _ruleset()
        hits = ruleset.search("", section="30", limit=500)
        self.assertEqual(len(hits), 157)
        self.assertTrue(all(h["sec"] == "30" for h in hits))

    def test_search_limit(self):
        ruleset = _ruleset()
        self.assertEqual(len(ruleset.search("a", limit=7)), 7)


class TestCatalogoService(ARQGenTestCase):
    """Fachada del servicio sobre el catálogo completo."""

    def test_service_search_and_official_price(self):
        from services.precons_service import PreconsService
        ctx = self.application.create_project(self.project_path(),
                                              name="CAT")
        service = PreconsService(ctx)
        service.load_ruleset(RULESET)
        self.assertEqual(len(service.sections()), 38)
        hits = service.search_items("manigua", limit=10)
        self.assertTrue(hits)
        self.assertIn("MANIGUA", hits[0]["name"].upper())
        price = service.official_price("011112")
        self.assertAlmostEqual(price["total"], 14.42, places=2)
        ctx.close()

    def test_service_requires_ruleset(self):
        from services.precons_service import PreconsService
        ctx = self.application.create_project(self.project_path(),
                                              name="CAT2")
        service = PreconsService(ctx)
        with self.assertRaises(DomainError):
            service.official_price("011111")
        ctx.close()


class TestCatalogoCLI(ARQGenTestCase):
    """CLI: listado filtrado y precio oficial (comando `precons`)."""

    def setUp(self):
        super().setUp()
        self.path = self.project_path()
        self.log_dir = self.tmp

    def cli(self, argv):
        import io
        from unittest.mock import patch
        from app.cli import run
        with patch("sys.stdout", new=io.StringIO()) as out:
            code = run(["--log-dir", self.log_dir] + argv)
        return code, out.getvalue()

    def _create_project(self):
        code, out = self.cli(["project", "create", self.path, "--name",
                              "Catálogo"])
        assert code == 0, out

    def test_list_filtered_by_search_and_section(self):
        self._create_project()
        code, out = self.cli(["precons", "list", self.path,
                              "--ruleset", RULESET,
                              "--buscar", "manigua", "--limite", "5"])
        self.assertEqual(code, 0, out)
        self.assertIn("coincidencias", out)
        self.assertIn("MANIGUA", out)

    def test_list_default_is_capped(self):
        self._create_project()
        code, out = self.cli(["precons", "list", self.path,
                              "--ruleset", RULESET])
        self.assertEqual(code, 0, out)
        self.assertIn("15981 partidas", out)
        self.assertIn("--limite 0", out)

    def test_official_price_command(self):
        self._create_project()
        code, out = self.cli(["precons", "precio", self.path,
                              "--ruleset", RULESET,
                              "--item", "011111", "--quantity", "100"])
        self.assertEqual(code, 0, out)
        self.assertIn("011111", out)
        self.assertIn("26.81", out)
        self.assertIn("2681.00", out)
        self.assertIn("CUP", out)

    def test_precio_unknown_item_fails_clean(self):
        self._create_project()
        code, out = self.cli(["precons", "precio", self.path,
                              "--ruleset", RULESET, "--item", "NO-EXISTE"])
        self.assertNotEqual(code, 0, out)


class TestImportadorPreconsIII(ARQGenTestCase):
    """Importador del libro oficial PRECONS III (shape "precons", spec 94)."""

    XLSX_REAL = os.path.normpath(os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        os.pardir, os.pardir, os.pardir,
        "catalogo_inbox", "Catalogo_PRECONS_III_Completo.xlsx"))

    def _make_libro(self, name: str = "libro_precons.xlsx") -> str:
        """Libro mínimo con las 9 hojas del formato PRECONS III."""
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        portada = wb.create_sheet("PORTADA")
        portada.append(["CATÁLOGO COMPLETO PRECONS III"])
        portada.append([None, "Versión de RoPres", "3.30"])
        portada.append([None, "Fecha de generación", "2026-09-24"])
        portada.append([None, "Renglones de listas de costos", 2])
        portada.append([None, "Materiales", 1])
        portada.append([None, "Equipos", 1])
        portada.append([None, "Mano de obra", 1])
        secciones = wb.create_sheet("SECCIONES")
        secciones.append(["Código", "Sección", "Renglones"])
        secciones.append(["01", "Seccion 01 - Prueba", 2])
        renglones = wb.create_sheet("RENGLONES_PRECONS")
        renglones.append(["Sección código", "Sección", "Capítulo código",
                          "Capítulo", "Subcapítulo código", "Subcapítulo",
                          "Código PRECONS", "Unidad", "Materiales CUP",
                          "Mano de obra CUP", "Equipos CUP", "Total CUP",
                          "Descripción"])
        renglones.append(["01", "Seccion 01 - Prueba", "011", "PRELIMINARES",
                          "0111", "CHAPEA", "011111", "m2", "26.81", "0",
                          "0", "26.81", "MANIGUA MANUAL"])
        renglones.append(["01", "Seccion 01 - Prueba", "011", "PRELIMINARES",
                          "0111", "CHAPEA", "011112", "m2", "0.66", "13.76",
                          "0", "14.42", "MANIGUA LIGERA"])
        materiales = wb.create_sheet("MATERIALES")
        materiales.append(["Código", "Descripción", "Unidad",
                           "Precio oficial CUP", "Peso kg"])
        materiales.append(["0001", "CEMENTO", "saco", "185", "42.5"])
        equipos = wb.create_sheet("EQUIPOS")
        equipos.append(["Código", "Descripción", "Unidad",
                        "Tarifa horaria CUP"])
        equipos.append(["E-01", "HORMIGONERA", "h", "332.65"])
        mano = wb.create_sheet("MANO_OBRA")
        mano.append(["Código", "Descripción", "Unidad", "Tarifa horaria CUP"])
        mano.append(["M-01", "ALBAÑIL", "h", "13.85"])
        params = wb.create_sheet("PARAMETROS")
        params.append(["Código", "Parámetro", "Valor", "Unidad", "Nota"])
        params.append(["horas_jornada", "Horas por jornada", "8", "h", ""])
        limites = wb.create_sheet("LIMITES")
        limites.append(["Código", "Límite / coeficiente", "Valor", "Unidad",
                        "Referencia"])
        limites.append(["UTILIDAD_MAX", "Utilidad máxima", "0.15", "%", "Hasta 15 %"])
        instr = wb.create_sheet("INSTRUCCIONES")
        instr.append(["Tema", "Detalle"])
        instr.append(["Alcance", "Libro de prueba."])
        path = os.path.join(self.tmp, name)
        wb.save(path)
        return path

    def test_shape_precons_end_to_end(self):
        path = self._make_libro()
        from importers.catalog import CatalogImporter
        importer = CatalogImporter(mapping={"code": "precons_test"})
        payload, report, prices = importer.convert(path)
        self.assertEqual(report["shape"], "precons")
        self.assertEqual(report["errors"], [])
        self.assertEqual(report["counts"]["work_items"], 2)
        self.assertEqual(report["counts"]["resources"], 3)
        self.assertEqual(report["counts"]["prices"], 3)
        # jerarquía + desglose oficial
        item = payload["work_items"][0]
        self.assertEqual(item["code"], "011111")
        self.assertEqual(item["sec"], "01")
        self.assertEqual(item["cap"], "011")
        self.assertEqual(item["sub"], "0111")
        self.assertAlmostEqual(item["price_ref"], 26.81, places=2)
        self.assertAlmostEqual(item["price_material"], 26.81, places=2)
        # recursos por tipo con precio y peso
        kinds = {r["code"]: r for r in payload["resources"]}
        self.assertEqual(kinds["0001"]["kind"], "MATERIAL")
        self.assertAlmostEqual(kinds["0001"]["price_ref"], 185.0, places=2)
        self.assertAlmostEqual(kinds["0001"]["peso_kg"], 42.5, places=2)
        self.assertEqual(kinds["E-01"]["kind"], "EQUIPMENT")
        self.assertEqual(kinds["M-01"]["kind"], "LABOR")
        # metadatos de PORTADA
        self.assertEqual(payload["version"], "3.30")
        self.assertEqual(payload["generated"], "2026-09-24")
        self.assertEqual(payload["parameters"]["horas_jornada"]["value"], 8)
        self.assertEqual(payload["limits"]["UTILIDAD_MAX"]["value"], 0.15)
        self.assertEqual(payload["limits"]["UTILIDAD_MAX"]["reference"],
                         "Hasta 15 %")
        self.assertEqual(payload["source"]["instrucciones"][0]["tema"],
                         "Alcance")
        self.assertEqual(payload["variants"][0]["code"], "BASE")
        # el payload carga como ruleset válido
        from engines.precons_engine import PreconsRuleset, official_price
        ruleset = PreconsRuleset.from_dict(payload)
        price = official_price(ruleset, "011112")
        self.assertAlmostEqual(price["total"], 14.42, places=2)
        # preview de precios para la lista del proyecto
        self.assertEqual(prices["count"], 3)

    def test_apply_writes_ruleset(self):
        path = self._make_libro()
        from importers.catalog import CatalogImporter
        out = os.path.join(self.tmp, "ruleset_salida.json")
        importer = CatalogImporter(mapping={"code": "precons_test"})
        payload, report = importer.apply(
            path, code="precons_test", name="Prueba PRECONS III",
            out_path=out)
        self.assertTrue(report["applied"])
        self.assertTrue(os.path.exists(out))
        from engines.precons_engine import PreconsRuleset
        ruleset = PreconsRuleset.from_file(out)
        self.assertEqual(ruleset.code, "precons_test")
        self.assertEqual(ruleset.version, "3.30")
        self.assertEqual(len(ruleset.work_items), 2)

    def test_libro_real_si_presente(self):
        real = os.path.normpath(self.XLSX_REAL)
        if not os.path.exists(real):
            self.skipTest("XLSX oficial no presente en este entorno")
        from importers.catalog import CatalogImporter
        importer = CatalogImporter(
            mapping={"code": "precons_iii_rospres_330"})
        payload, report, _ = importer.convert(real)
        self.assertEqual(report["shape"], "precons")
        self.assertEqual(report["errors"], [])
        self.assertEqual(report["counts"]["work_items"], 15981)
        self.assertEqual(report["counts"]["resources"], 4383)
        self.assertEqual(payload["version"], "3.30")
        self.assertEqual(payload["generated"], "2026-09-24")
        self.assertEqual(payload["source"]["rospres"], "3.30")


class TestCatalogoGUI(ARQGenTestCase):
    """Navegador del catálogo (Herramientas ▸ Catálogo PRECONS III)."""

    def setUp(self):
        super().setUp()
        try:
            from ui.app_window import ARQGenWindow
            self.context = self.demo_context()
            self.window = ARQGenWindow(self.context, self.application)
            self.window.update()
        except Exception:  # noqa: BLE001 — sin display
            self.window = None
            self.skipTest("sin display para pruebas GUI")

    def tearDown(self):
        if self.window is not None:
            self.window.destroy()
            self.context.close()
        super().tearDown()

    def _pump(self, times=25):
        for _ in range(times):
            self.window.update()

    def _wait_loaded(self, timeout_s=20) -> None:
        """Bombea la cola Tk hasta que el diálogo cargó el catálogo."""
        import time
        deadline = time.time() + timeout_s
        while time.time() < deadline:
            self.window.update()
            if self.window._cat_section["values"]:
                return
            time.sleep(0.05)
        self.fail("el diálogo del catálogo no terminó de cargar")

    def test_catalog_browser_dialog(self):
        self.window.open_catalogo_precons()
        self._wait_loaded()
        dialog = self.window._catalogo_dialog
        self.assertTrue(dialog.winfo_exists())
        values = self.window._cat_section["values"]
        self.assertTrue(values)
        self.assertIn("01 — Seccion 01", values[1])
        # búsqueda por texto
        self.window._cat_search.delete(0, "end")
        self.window._cat_search.insert(0, "manigua")
        self.window._cat_run_search()
        self._pump(5)
        children = self.window._cat_tree.get_children()
        self.assertTrue(children)
        self.assertIn("coincidencias", self.window._cat_count.get())
        # selección → desglose + total por cantidad
        first = children[0]
        self.window._cat_tree.selection_set(first)
        self.window._cat_tree.event_generate("<<TreeviewSelect>>")
        self._pump(5)
        self.assertIn("Materiales", self.window._cat_detail.get())
        self.window._cat_qty.set("100")
        self.window._cat_show_detail()
        self._pump(5)
        self.assertIn("CUP", self.window._cat_total.get())
        # sin resultados: mensaje amable
        self.window._cat_search.delete(0, "end")
        self.window._cat_search.insert(0, "zzz-nada")
        self.window._cat_run_search()
        self._pump(5)
        self.assertIn("Sin resultados", self.window._cat_detail.get())
        dialog.destroy()

    def test_catalog_browser_totals_match_engine(self):
        from engines.precons_engine import official_price
        self.window.open_catalogo_precons()
        self._wait_loaded()
        children = self.window._cat_tree.get_children()
        self.assertTrue(children)
        first = children[0]
        self.window._cat_tree.selection_set(first)
        self.window._cat_tree.event_generate("<<TreeviewSelect>>")
        self._pump(5)
        code = first
        price = official_price(self.window._catalogo_ruleset(), code)
        self.assertAlmostEqual(price["total"],
                               float(self.window._catalogo_ruleset()
                                     .item(code)["price_ref"]), places=2)
        dialog = self.window._catalogo_dialog
        dialog.destroy()


if __name__ == "__main__":
    unittest.main()
